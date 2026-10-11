"""Excel/CSV import execution + bulk actions.

The frontend handles upload -> column mapping -> preview; this service runs
the validated import row-by-row, collecting per-row errors so the user can
download an error report. Also implements bulk activate/inactivate, bulk GST
rate update and bulk category update per the spec.
"""
import csv
import io
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional
from uuid import UUID

from sqlalchemy.orm import Session

from app.models.ledger import Ledger
from app.models.ledger_group import LedgerGroup
from app.models.party import Party
from app.models.product_item import ProductItem
from app.schemas.ledger import LedgerCreate
from app.schemas.party import PartyCreate, NameContactInfo, AdditionalInfo, PaymentInfo
from app.schemas.product_item import ProductItemCreate
from app.services.ledger_service import LedgerService
from app.services.party_service import PartyService
from app.services.product_item_service import ProductItemService
from app.services.gst_service import GstService


def _dec(v, default=None):
    if v is None or v == "":
        return default
    try:
        return Decimal(str(v).replace(",", ""))
    except (InvalidOperation, ValueError):
        raise ValueError(f"'{v}' is not a valid number")


# System target fields per import entity (used for auto column mapping).
ENTITY_FIELDS: Dict[str, List[str]] = {
    "LEDGER": ["name", "group", "gstin", "balance_type", "opening_balance", "description"],
    "CUSTOMER": ["name", "gstin", "pan", "company", "email", "mobile", "terms",
                 "credit_limit", "opening_balance"],
    "VENDOR": ["name", "gstin", "pan", "company", "email", "mobile", "terms",
               "credit_limit", "opening_balance"],
    "ITEM": ["item_code", "item_name", "item_type", "hsn_sac", "unit",
             "purchase_rate", "sale_rate", "gst_percent", "opening_stock"],
}


def _norm(s: str) -> str:
    return "".join(ch for ch in str(s or "").lower() if ch.isalnum())


class DataImportService:
    @staticmethod
    def parse_file(filename: str, content: bytes) -> Dict[str, Any]:
        """Parse an uploaded CSV or Excel (.xlsx) file into columns + rows."""
        name = (filename or "").lower()
        columns: List[str] = []
        rows: List[Dict[str, Any]] = []
        if name.endswith((".xlsx", ".xlsm")):
            try:
                from openpyxl import load_workbook
            except ImportError:
                raise ValueError("Excel support is unavailable on this server")
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            ws = wb.active
            it = ws.iter_rows(values_only=True)
            try:
                header = next(it)
            except StopIteration:
                header = ()
            columns = [str(h).strip() if h is not None else f"col{i+1}"
                       for i, h in enumerate(header)]
            for r in it:
                if r is None or all(c is None for c in r):
                    continue
                row = {}
                for i, col in enumerate(columns):
                    val = r[i] if i < len(r) else None
                    row[col] = "" if val is None else val
                rows.append(row)
        else:
            text = content.decode("utf-8-sig", errors="replace")
            reader = csv.DictReader(io.StringIO(text))
            columns = [c.strip() for c in (reader.fieldnames or [])]
            for r in reader:
                rows.append({(k.strip() if k else k): (v or "") for k, v in r.items()})
        return {"columns": columns, "rows": rows, "row_count": len(rows)}

    @staticmethod
    def suggest_mapping(entity: str, columns: List[str]) -> Dict[str, Optional[str]]:
        """Auto-map source columns to system fields for the given entity."""
        entity = (entity or "").upper()
        fields = ENTITY_FIELDS.get(entity, [])
        norm_cols = {_norm(c): c for c in columns}
        mapping: Dict[str, Optional[str]] = {}
        aliases = {
            "name": ["name", "ledgername", "partyname", "customername", "vendorname",
                     "displayname", "accountname"],
            "group": ["group", "ledgergroup", "groupname", "category"],
            "gstin": ["gstin", "gst", "gstno", "gstnumber"],
            "pan": ["pan", "panno", "pannumber"],
            "company": ["company", "companyname", "firm"],
            "email": ["email", "emailid", "mail"],
            "mobile": ["mobile", "phone", "mobilenumber", "contact", "mobileno"],
            "terms": ["terms", "paymentterms", "creditterms"],
            "credit_limit": ["creditlimit", "credit"],
            "opening_balance": ["openingbalance", "opening", "openingbal", "balance"],
            "balance_type": ["balancetype", "drcr", "type"],
            "description": ["description", "notes", "remarks"],
            "item_code": ["itemcode", "code", "sku"],
            "item_name": ["itemname", "name", "product", "productname"],
            "item_type": ["itemtype", "type", "goodsservice"],
            "hsn_sac": ["hsnsac", "hsn", "sac", "hsncode"],
            "unit": ["unit", "uom", "units"],
            "purchase_rate": ["purchaserate", "purchaseprice", "costprice", "buyrate"],
            "sale_rate": ["salerate", "saleprice", "sellingprice", "mrp", "rate"],
            "gst_percent": ["gstpercent", "gstrate", "taxrate", "gst"],
            "opening_stock": ["openingstock", "stock", "qty", "quantity"],
        }
        for field in fields:
            found = None
            for alias in aliases.get(field, [field]):
                if alias in norm_cols:
                    found = norm_cols[alias]
                    break
            if not found:
                # fuzzy: source column contains the field token
                for nc, orig in norm_cols.items():
                    if _norm(field) and _norm(field) in nc:
                        found = orig
                        break
            mapping[field] = found
        return mapping

    @staticmethod
    def apply_mapping(
        rows: List[Dict[str, Any]], mapping: Dict[str, Optional[str]]
    ) -> List[Dict[str, Any]]:
        """Rename source columns to system field names using the mapping."""
        out: List[Dict[str, Any]] = []
        for r in rows:
            mapped: Dict[str, Any] = {}
            for field, source_col in mapping.items():
                if source_col is not None and source_col in r:
                    mapped[field] = r[source_col]
            out.append(mapped)
        return out

    @staticmethod
    def execute(db: Session, entity: str, rows: List[Dict[str, Any]]) -> Dict[str, Any]:
        entity = (entity or "").upper()
        handler = {
            "LEDGER": DataImportService._import_ledger,
            "CUSTOMER": lambda db, r: DataImportService._import_party(db, r, "C"),
            "VENDOR": lambda db, r: DataImportService._import_party(db, r, "V"),
            "ITEM": DataImportService._import_item,
        }.get(entity)
        if not handler:
            return {"imported": 0, "failed": len(rows), "errors": [
                {"row": 0, "message": f"Unknown entity {entity}"}]}

        imported = 0
        errors: List[Dict[str, Any]] = []
        for idx, row in enumerate(rows, start=1):
            try:
                handler(db, row)
                imported += 1
            except Exception as e:  # collect and continue
                db.rollback()
                errors.append({"row": idx, "data": row, "message": str(e)})
        return {"imported": imported, "failed": len(errors), "errors": errors}

    @staticmethod
    def _import_ledger(db: Session, row: Dict[str, Any]) -> None:
        name = (row.get("name") or "").strip()
        if not name:
            raise ValueError("Ledger name is required")
        group_id: Optional[UUID] = None
        group_name = (row.get("group") or "").strip()
        if group_name:
            grp = db.query(LedgerGroup).filter(
                LedgerGroup.name.ilike(group_name), LedgerGroup.active == "Y"
            ).first()
            if not grp:
                raise ValueError(f"Ledger group '{group_name}' not found")
            group_id = grp.id
        gstin = (row.get("gstin") or "").strip().upper() or None
        if gstin:
            v = GstService.validate_gstin(gstin)
            if v["level"] == "error":
                raise ValueError(v["message"])
        payload = LedgerCreate(
            name=name,
            ledger_group_id=group_id,
            balance_type=(row.get("balance_type") or "").strip().upper() or None,
            opening_balance=_dec(row.get("opening_balance"), Decimal("0")),
            gstin=gstin,
            description=(row.get("description") or None),
            active="Y",
        )
        data, status_val, message = LedgerService.create(db, payload)
        if status_val == "ERROR":
            raise ValueError(message)

    @staticmethod
    def _import_party(db: Session, row: Dict[str, Any], key_group: str) -> None:
        name = (row.get("name") or "").strip()
        if not name:
            raise ValueError("Party name is required")
        if PartyService.find_duplicate_name(db, name, key_group, None):
            raise ValueError(f"Party '{name}' already exists")
        gstin = (row.get("gstin") or "").strip().upper() or None
        if gstin:
            v = GstService.validate_gstin(gstin)
            if v["level"] == "error":
                raise ValueError(v["message"])
        payload = PartyCreate(
            key_group=key_group,
            active="Y",
            gstin=gstin,
            pan=(row.get("pan") or None),
            name_contact_info=NameContactInfo(
                first_name=name,
                display_name=name,
                company_name=(row.get("company") or None),
                email=(row.get("email") or None),
                mobile_number=(str(row.get("mobile")) if row.get("mobile") else None),
            ),
            payment_info=PaymentInfo(
                terms=(row.get("terms") or None),
                credit_limit=_dec(row.get("credit_limit"), Decimal("0")),
            ),
            additional_info=AdditionalInfo(
                opening_bal=_dec(row.get("opening_balance"), Decimal("0")),
            ),
        )
        data, status_val, message = PartyService.create(db, payload)
        if status_val == "ERROR":
            raise ValueError(message)

    @staticmethod
    def _import_item(db: Session, row: Dict[str, Any]) -> None:
        name = (row.get("item_name") or row.get("name") or "").strip()
        if not name:
            raise ValueError("Item name is required")
        code = (row.get("item_code") or "").strip() or ProductItemService.next_code(db)
        payload = ProductItemCreate(
            item_code=code,
            item_name=name,
            item_type=(row.get("item_type") or "G")[:1].upper(),
            hsn_sac=(str(row.get("hsn_sac")) if row.get("hsn_sac") else None),
            unit=(row.get("unit") or "PCS"),
            purchase_rate=_dec(row.get("purchase_rate"), Decimal("0")),
            sale_rate=_dec(row.get("sale_rate"), Decimal("0")),
            gst_percent=_dec(row.get("gst_percent"), Decimal("0")),
            opening_stock=_dec(row.get("opening_stock"), Decimal("0")),
            status="Active",
        )
        ProductItemService.create(db, payload)


class BulkActionService:
    @staticmethod
    def bulk_status(db: Session, entity: str, ids: List[str], active: str) -> Dict[str, Any]:
        entity = (entity or "").upper()
        count = 0
        if entity == "ITEM":
            status = "Active" if active == "Y" else "Inactive"
            for iid in ids:
                item = db.query(ProductItem).filter(ProductItem.id == int(iid)).first()
                if item:
                    item.status = status
                    count += 1
        elif entity == "LEDGER":
            for lid in ids:
                l = db.query(Ledger).filter(Ledger.id == UUID(lid)).first()
                if l:
                    l.active = active
                    count += 1
        elif entity in ("CUSTOMER", "VENDOR", "PARTY"):
            for pid in ids:
                p = db.query(Party).filter(Party.id == UUID(pid)).first()
                if p:
                    p.active = active
                    if p.ledger_id:
                        l = db.query(Ledger).filter(Ledger.id == p.ledger_id).first()
                        if l:
                            l.active = active
                    count += 1
        db.commit()
        return {"updated": count}

    @staticmethod
    def bulk_delete(db: Session, entity: str, ids: List[str]) -> Dict[str, Any]:
        """Restricted bulk delete: soft-delete only, skipping system/frozen records."""
        entity = (entity or "").upper()
        deleted = 0
        skipped: List[str] = []
        if entity == "ITEM":
            for iid in ids:
                item = db.query(ProductItem).filter(ProductItem.id == int(iid)).first()
                if item:
                    item.status = "Inactive"
                    deleted += 1
        elif entity == "LEDGER":
            for lid in ids:
                l = db.query(Ledger).filter(Ledger.id == UUID(lid)).first()
                if not l:
                    continue
                if l.is_system == "Y":
                    skipped.append(l.name + " (system)")
                    continue
                if l.is_frozen == "Y":
                    skipped.append(l.name + " (frozen)")
                    continue
                l.active = "N"
                deleted += 1
        elif entity in ("CUSTOMER", "VENDOR", "PARTY"):
            for pid in ids:
                p = db.query(Party).filter(Party.id == UUID(pid)).first()
                if p:
                    p.active = "N"
                    if p.ledger_id:
                        l = db.query(Ledger).filter(Ledger.id == p.ledger_id).first()
                        if l:
                            l.active = "N"
                    deleted += 1
        db.commit()
        return {"deleted": deleted, "skipped": skipped}

    @staticmethod
    def bulk_gst_rate(db: Session, ids: List[str], gst_percent) -> Dict[str, Any]:
        count = 0
        for iid in ids:
            item = db.query(ProductItem).filter(ProductItem.id == int(iid)).first()
            if item:
                item.gst_percent = gst_percent
                count += 1
        db.commit()
        return {"updated": count}

    @staticmethod
    def bulk_category(db: Session, ids: List[str], ledger_group_id: str) -> Dict[str, Any]:
        gid = UUID(ledger_group_id)
        count = 0
        for lid in ids:
            l = db.query(Ledger).filter(Ledger.id == UUID(lid)).first()
            if l:
                l.ledger_group_id = gid
                count += 1
        db.commit()
        return {"updated": count}
